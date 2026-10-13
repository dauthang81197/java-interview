#!/usr/bin/env python3
"""Đánh dấu bài đã làm trong tracker/bai-tap-java.xlsx.

Dùng (sau khi Claude đã review bài):
    .venv/bin/python tracker/mark_done.py JR-01 --reviewed
    .venv/bin/python tracker/mark_done.py JR-01 MID-02 --reviewed
    .venv/bin/python tracker/mark_done.py JR-01 --undo        # đặt lại "Chưa làm"

Script CHỈ đánh dấu khi: (1) có cờ --reviewed (xác nhận Claude đã review),
(2) test của bài đó xanh (script tự chạy `mvn -q test -Dtest=<Test>`).
"""
import argparse
import subprocess
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = Path(__file__).with_name("bai-tap-java.xlsx")
PENDING, DONE = "Chưa làm", "Đã làm"


def columns(ws):
    """Tìm cột theo tiêu đề (không phụ thuộc vị trí)."""
    hdr = {c.value: c.column for c in ws[1] if c.value}
    need = ["Mã bài", "Test", "Trạng thái", "Ngày hoàn thành"]
    missing = [h for h in need if h not in hdr]
    if missing:
        raise SystemExit(f"Excel thiếu cột: {', '.join(missing)}. Chạy lại build_tracker.py")
    return hdr


def find_row(ws, code, col_code):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, col_code).value == code:
            return r
    return None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("codes", nargs="+", help="mã bài, ví dụ JR-01")
    ap.add_argument("--reviewed", action="store_true", help="Claude đã review bài này")
    ap.add_argument("--undo", action="store_true", help="đặt lại trạng thái Chưa làm")
    args = ap.parse_args()

    try:
        wb = load_workbook(XLSX)
    except FileNotFoundError:
        print(f"Không thấy {XLSX}. Chạy trước: .venv/bin/python tracker/build_tracker.py", file=sys.stderr)
        return 2
    ws = wb.active
    hdr = columns(ws)
    COL_TEST, COL_STATUS, COL_DATE = hdr["Test"], hdr["Trạng thái"], hdr["Ngày hoàn thành"]
    rc = 0
    changed = False
    for code in (c.upper() for c in args.codes):
        row = find_row(ws, code, hdr["Mã bài"])
        if row is None:
            print(f"[{code}] không có trong Excel", file=sys.stderr)
            rc = 1
            continue
        if args.undo:
            ws.cell(row, COL_STATUS).value = PENDING
            ws.cell(row, COL_DATE).value = None
            changed = True
            print(f"[{code}] đã đặt lại: {PENDING}")
            continue
        if not args.reviewed:
            print(f"[{code}] bỏ qua: thiếu --reviewed (cần Claude review trước)", file=sys.stderr)
            rc = 1
            continue
        test = ws.cell(row, COL_TEST).value
        print(f"[{code}] chạy test {test} ...")
        try:
            res = subprocess.run(["mvn", "-q", "test", f"-Dtest={test}"], cwd=ROOT,
                                 capture_output=True, text=True, timeout=600)
        except FileNotFoundError:
            print(f"[{code}] không tìm thấy lệnh mvn (cần cài Maven và có trong PATH).", file=sys.stderr)
            rc = 1
            continue
        except subprocess.TimeoutExpired:
            print(f"[{code}] mvn chạy quá 600s, coi như test chưa xanh.", file=sys.stderr)
            rc = 1
            continue
        if res.returncode != 0:
            print(f"[{code}] test CHƯA xanh, không đánh dấu.\n" + "\n".join(res.stdout.splitlines()[-25:]),
                  file=sys.stderr)
            rc = 1
            continue
        ws.cell(row, COL_STATUS).value = DONE
        ws.cell(row, COL_DATE).value = date.today().isoformat()
        changed = True
        print(f"[{code}] -> {DONE}")
    if changed:
        try:
            wb.save(XLSX)
        except PermissionError:
            print("Không ghi được file Excel (đang mở trong Excel?). Đóng file rồi chạy lại.", file=sys.stderr)
            return 2
    return rc


if __name__ == "__main__":
    sys.exit(main())
