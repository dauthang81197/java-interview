#!/usr/bin/env python3
"""Sinh tracker/bai-tap-java.xlsx từ danh sách bài tập.

Chạy: .venv/bin/python tracker/build_tracker.py
Cẩn thận: ghi đè file Excel; nếu file đã tồn tại sẽ giữ lại trạng thái các bài đã "Đã làm".
Khi thêm bài mới: thêm các dòng vào EXERCISES rồi chạy lại.
"""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["STT", "Mã bài", "Cấp độ", "Chủ đề", "Tên bài", "File bài tập", "Test", "Độ khó",
           "Mức hay hỏi", "Trạng thái", "Ngày hoàn thành", "Ghi chú"]
PENDING = "Chưa làm"
DONE = "Đã làm"

M = "src/main/java/com/example/javaex"
# (STT, mã, cấp độ, chủ đề, tên, file bài tập (không .java), độ khó, mức hay hỏi)
# Tên class test = <tên class bài tập> + "Test".
EXERCISES = [
    (1, "JR-01", "Junior", "Cơ bản", "Kiểu nguyên thủy, tràn số, ép kiểu", "junior/t01_coban/Bai01JrPrimitives", "Dễ", "Cao"),
    (2, "JR-02", "Junior", "Cơ bản", "==, equals, autoboxing, Integer cache", "junior/t01_coban/Bai02JrBoxing", "Dễ", "Rất cao"),
    (3, "JR-03", "Junior", "Cơ bản", "String bất biến và StringBuilder", "junior/t01_coban/Bai03JrStringBuilder", "Dễ", "Rất cao"),
    (4, "JR-04", "Junior", "Cơ bản", "Truyền tham số: pass-by-value", "junior/t01_coban/Bai04JrPassByValue", "Dễ", "Cao"),
    (5, "JR-05", "Junior", "Cơ bản", "Enum, switch expression và record", "junior/t01_coban/Bai05JrEnumRecord", "Dễ", "Trung bình"),
    (6, "JR-06", "Junior", "OOP", "Kế thừa, đa hình, override và overload", "junior/t02_oop/Bai06JrPolymorphism", "Dễ", "Rất cao"),
    (7, "JR-07", "Junior", "OOP", "Abstract class và interface", "junior/t02_oop/Bai07JrInterfaces", "Dễ", "Rất cao"),
    (8, "JR-08", "Junior", "OOP", "Thứ tự khởi tạo: static, instance, constructor", "junior/t02_oop/Bai08JrInitOrder", "Trung bình", "Cao"),
    (9, "JR-09", "Junior", "Exception", "Exception: checked, unchecked, finally, try-with-resources", "junior/t03_exception/Bai09JrExceptions", "Dễ", "Rất cao"),
    (10, "JR-10", "Junior", "Collections và Generics", "Collections cơ bản: List/Set/Map, Comparable", "junior/t04_collections/Bai10JrCollectionsBasics", "Dễ", "Rất cao"),
    (11, "JR-11", "Junior", "Collections và Generics", "Generics cơ bản", "junior/t04_collections/Bai11JrGenerics", "Dễ", "Cao"),
    (12, "JR-12", "Junior", "Lambda và Optional", "Lambda và functional interface cơ bản", "junior/t05_functional/Bai12JrLambdas", "Dễ", "Cao"),
    (13, "JR-13", "Junior", "Lambda và Optional", "Optional cơ bản", "junior/t05_functional/Bai13JrOptional", "Dễ", "Trung bình"),
    (14, "JR-14", "Junior", "Thread và đồng bộ cơ bản", "Thread: start() vs run(), vòng đời, daemon", "junior/t06_concurrency/Bai14JrThreadBasics", "Dễ", "Rất cao"),
    (15, "JR-15", "Junior", "Thread và đồng bộ cơ bản", "Race condition và synchronized", "junior/t06_concurrency/Bai15JrRaceCondition", "Dễ", "Rất cao"),
    (16, "MID-01", "Mid", "Nền tảng nâng cao", "String pool và so sánh tham chiếu", "mid/t01_core/Bai16MidStringPoolQuiz", "Dễ", "Cao"),
    (17, "MID-02", "Mid", "Nền tảng nâng cao", "equals/hashCode đúng hợp đồng", "mid/t01_core/Bai17MidMoney", "Dễ", "Rất cao"),
    (18, "MID-03", "Mid", "Nền tảng nâng cao", "Thiết kế lớp bất biến", "mid/t01_core/Bai18MidImmutableBooking", "Dễ", "Cao"),
    (19, "MID-04", "Mid", "Nền tảng nâng cao", "Inner class, anonymous class và lambda capture", "mid/t01_core/Bai19MidClosures", "Trung bình", "Cao"),
    (20, "MID-05", "Mid", "Exception", "Exception chaining, custom exception, suppressed", "mid/t02_exception/Bai20MidExceptionChaining", "Trung bình", "Cao"),
    (21, "MID-06", "Mid", "Collections, Stream và Generics", "Comparator chaining và sắp xếp ổn định", "mid/t03_collections/Bai21MidComparators", "Trung bình", "Cao"),
    (22, "MID-07", "Mid", "Collections, Stream và Generics", "ArrayList tự viết + iterator fail-fast", "mid/t03_collections/Bai22MidMyArrayList", "Trung bình", "Cao"),
    (23, "MID-08", "Mid", "Collections, Stream và Generics", "Generics và PECS", "mid/t03_collections/Bai23MidGenericsExercises", "Khó", "Cao"),
    (24, "MID-09", "Mid", "Collections, Stream và Generics", "Stream và Collectors", "mid/t03_collections/Bai24MidStreamExercises", "Trung bình", "Rất cao"),
    (25, "MID-10", "Mid", "Collections, Stream và Generics", "LRU Cache", "mid/t03_collections/Bai25MidLruCache", "Trung bình", "Rất cao"),
    (26, "MID-11", "Mid", "Java 21 hiện đại", "Sealed, record và pattern matching (Java 21)", "mid/t04_modern/Bai26MidSealedPatterns", "Trung bình", "Trung bình"),
    (27, "MID-12", "Mid", "Concurrency cơ bản", "volatile, visibility và happens-before", "mid/t05_concurrency/Bai27MidVolatileFlag", "Trung bình", "Rất cao"),
    (28, "MID-13", "Mid", "Concurrency cơ bản", "Singleton lazy thread-safe", "mid/t05_concurrency/Bai28MidLazyConfig", "Dễ", "Rất cao"),
    (29, "MID-14", "Mid", "Concurrency cơ bản", "wait/notify: hộp thư một ô (producer-consumer đơn giản)", "mid/t05_concurrency/Bai29MidWaitNotify", "Trung bình", "Cao"),
    (30, "MID-15", "Mid", "Concurrency cơ bản", "Atomic, CAS và lock-free", "mid/t05_concurrency/Bai30MidAtomic", "Trung bình", "Cao"),
    (31, "MID-16", "Mid", "Concurrency cơ bản", "ExecutorService cơ bản", "mid/t05_concurrency/Bai31MidExecutor", "Trung bình", "Rất cao"),
    (32, "SR-01", "Senior", "JVM và bộ nhớ", "Stack không rò rỉ bộ nhớ", "senior/t01_jvm/Bai32SrLeakyStack", "Dễ", "Cao"),
    (33, "SR-02", "Senior", "JVM và bộ nhớ", "Quan sát heap và GC", "senior/t01_jvm/Bai33SrJvmInfo", "Dễ", "Rất cao"),
    (34, "SR-03", "Senior", "JVM và bộ nhớ", "Cache giá trị tham chiếu yếu", "senior/t01_jvm/Bai34SrWeakValueCache", "Trung bình", "Trung bình"),
    (35, "SR-04", "Senior", "JVM và bộ nhớ", "ClassLoader tự viết", "senior/t01_jvm/Bai35SrByteArrayClassLoader", "Khó", "Cao"),
    (36, "SR-05", "Senior", "Collections nội bộ", "HashMap tự viết", "senior/t02_collections/Bai36SrSimpleHashMap", "Khó", "Rất cao"),
    (37, "SR-06", "Senior", "Đa luồng", "Phát hiện deadlock bằng ThreadMXBean", "senior/t03_concurrency/Bai37SrDeadlockDetector", "Trung bình", "Cao"),
    (38, "SR-07", "Senior", "Đa luồng", "Ping-pong hai luồng", "senior/t03_concurrency/Bai38SrPingPong", "Trung bình", "Cao"),
    (39, "SR-08", "Senior", "Đa luồng", "Blocking queue (producer-consumer)", "senior/t03_concurrency/Bai39SrBoundedBlockingQueue", "Khó", "Rất cao"),
    (40, "SR-09", "Senior", "Đa luồng", "Chuyển tiền không deadlock", "senior/t03_concurrency/Bai40SrBank", "Khó", "Rất cao"),
    (41, "SR-10", "Senior", "Đa luồng", "Tổng song song: Executor và Fork/Join", "senior/t03_concurrency/Bai41SrParallelSum", "Trung bình", "Cao"),
    (42, "SR-11", "Senior", "Đa luồng", "Memoizer thread-safe", "senior/t03_concurrency/Bai42SrMemoizer", "Khó", "Cao"),
    (43, "SR-12", "Senior", "Đa luồng", "CompletableFuture", "senior/t03_concurrency/Bai43SrAsyncAggregator", "Khó", "Cao"),
]


def main() -> None:
    out = Path(__file__).with_name("bai-tap-java.xlsx")
    old = {}
    if out.exists():
        ws_old = load_workbook(out).active
        hdr = {c.value: i for i, c in enumerate(ws_old[1]) if c.value}
        for row in ws_old.iter_rows(min_row=2, values_only=True):
            code = row[hdr["Mã bài"]]
            if code:
                old[code] = (row[hdr["Trạng thái"]], row[hdr["Ngày hoàn thành"]], row[hdr["Ghi chú"]])

    wb = Workbook()
    ws = wb.active
    ws.title = "Bài tập"
    ws.append(HEADERS)
    for stt, code, level, topic, name, path, diff, freq in EXERCISES:
        status, date, note = old.get(code, (PENDING, None, None))
        test = path.split("/")[-1] + "Test"
        ws.append([stt, code, level, topic, name, f"{M}/{path}.java", test, diff, freq, status, date, note])

    head_fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for col, w in zip("ABCDEFGHIJKL", [6, 9, 10, 14, 44, 66, 38, 12, 14, 12, 18, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    dv = DataValidation(type="list", formula1=f'"{PENDING},{DONE}"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{ws.max_row}")

    n = len(EXERCISES)
    ws["N1"], ws["O1"] = "Tổng số bài", n
    ws["N2"], ws["O2"] = "Đã làm", f'=COUNTIF(J2:J{n + 1},"{DONE}")'
    ws["N3"], ws["O3"] = "Tiến độ", "=O2/O1"
    ws["O3"].number_format = "0%"
    for r, lv in zip((5, 6, 7), ("Junior", "Mid", "Senior")):
        ws[f"N{r}"] = lv
        ws[f"O{r}"] = f'=COUNTIFS(C2:C{n + 1},"{lv}",J2:J{n + 1},"{DONE}")&"/"&COUNTIF(C2:C{n + 1},"{lv}")'
    for r in (1, 2, 3, 5, 6, 7):
        ws[f"N{r}"].font = Font(bold=True)
    ws.column_dimensions["N"].width = 14

    wb.save(out)
    print(f"Đã ghi {out} ({n} bài)")


if __name__ == "__main__":
    main()
